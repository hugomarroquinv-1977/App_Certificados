# app.py — LA COCINA de Academia Horizonte.
# Toda la lógica vive aquí. El HTML (salón) solo muestra lo ya calculado.
# Baldosas usadas abajo: variables, tipos, condicionales, bucles y funciones.
from collections import defaultdict  # VARIABLE: defaultdict es un tipo diccionario con valor por defecto.
from datetime import date  # VARIABLE: date sirve para poner la fecha de emisión.
from functools import wraps  # VARIABLE: wraps conserva el nombre de la función que decoro.
import io  # VARIABLE: io permite leer y escribir archivos en memoria, sin tocar el disco.
import os  # VARIABLE: os sirve para leer las variables de entorno (credenciales).
from pathlib import Path  # VARIABLE: Path es un tipo para rutas de archivos.
from flask import Flask, render_template, request, session, redirect, url_for, abort, send_file  # VARIABLE: session guarda quién entró; redirect y url_for mueven entre páginas; abort corta con un error HTTP; send_file devuelve un archivo.
from dotenv import load_dotenv  # VARIABLE: load_dotenv lee el archivo .env de la carpeta.
import openpyxl  # VARIABLE: openpyxl es la librería para leer y escribir los Excel.

# VARIABLE de tipo Path: carpeta donde vive app.py, para encontrar Insumos/ desde donde sea.
BASE_DIR = Path(__file__).resolve().parent
# FUNCIÓN: carga el archivo .env de esta carpeta (en Vercel ese archivo no existe y se usan las variables del panel).
load_dotenv(BASE_DIR / ".env")

# VARIABLES de tipo Path con las rutas de la bodega (solo lectura, NO modificar).
RUTA_MAESTRO = BASE_DIR / "Insumos" / "Maestro_Estudiantes.xlsx"
RUTA_EVALS = BASE_DIR / "Insumos" / "Registro_Evaluaciones.xlsx"

# VARIABLE de tipo str: texto que muestra de dónde salieron los datos en pantalla.
ORIGEN_FABRICA = "Datos de fábrica (carpeta Insumos/)"
# VARIABLE de tipo int: tope de la petición completa; Vercel corta las de más de 4.5 MB.
LIMITE_PETICION = int(4.5 * 1024 * 1024)
# VARIABLE de tipo int: tope por archivo, para avisar con un mensaje entendible.
LIMITE_ARCHIVO = 4 * 1024 * 1024


# FUNCIÓN: lee una variable de entorno y devuelve None si no existe o está vacía.
def leer_variable(nombre):
    # VARIABLE de tipo str o None: lo que tenga la variable en el entorno.
    valor = os.environ.get(nombre)
    # CONDICIONAL: si no existe, no hay nada que limpiar.
    if valor is None:
        return None
    valor = valor.strip()
    # CONDICIONAL: una variable vacía cuenta como no configurada.
    if valor == "":
        return None
    return valor


# VARIABLES de tipo str: credenciales del administrador y del usuario normal.
ADMIN_CORREO = leer_variable("ADMIN_CORREO")
ADMIN_CLAVE = leer_variable("ADMIN_CLAVE")
USUARIO_CORREO = leer_variable("USUARIO_CORREO")
USUARIO_CLAVE = leer_variable("USUARIO_CLAVE")

# VARIABLE de tipo bool: True si falta alguna credencial; entonces uso datos de demostración.
MODO_DEMO = (ADMIN_CORREO is None or ADMIN_CLAVE is None or
             USUARIO_CORREO is None or USUARIO_CLAVE is None)
# CONDICIONAL: completo solo lo que falte con valores de prueba (la pantalla avisa que es demostración).
if MODO_DEMO:
    ADMIN_CORREO = ADMIN_CORREO or "admin@ejemplo.cr"
    ADMIN_CLAVE = ADMIN_CLAVE or "admin123"
    USUARIO_CORREO = USUARIO_CORREO or "usuario@ejemplo.cr"
    USUARIO_CLAVE = USUARIO_CLAVE or "usuario123"

# VARIABLE de tipo dict: los dos usuarios válidos; la clave del diccionario es el correo.
USUARIOS = {
    ADMIN_CORREO: {"clave": ADMIN_CLAVE, "rol": "admin"},
    USUARIO_CORREO: {"clave": USUARIO_CLAVE, "rol": "usuario"},
}

# VARIABLE de tipo Flask: la aplicación web.
app = Flask(__name__)
# VARIABLE de tipo str: clave para firmar la sesión y que nadie pueda falsificarla.
app.secret_key = os.environ.get("SECRET_KEY") or "clave-demo-cambiar-en-produccion"
# VARIABLE de configuración: límite total de la petición; si se pasa, Flask responde 413.
app.config["MAX_CONTENT_LENGTH"] = LIMITE_PETICION


# FUNCIÓN: compara correo y clave contra la lista; devuelve el usuario o None.
def verificar_credenciales(correo, clave):
    # VARIABLE de tipo dict o None: el usuario que corresponde a ese correo.
    usuario = USUARIOS.get(correo)
    # CONDICIONAL: correo desconocido.
    if usuario is None:
        return None
    # CONDICIONAL: la clave no coincide.
    if usuario["clave"] != clave:
        return None
    return usuario


# FUNCIÓN decoradora: deja pasar solo si hay una sesión iniciada.
def requiere_login(vista):
    @wraps(vista)
    def envoltura(*args, **kwargs):
        # CONDICIONAL: si no hay correo guardado en la sesión, mando a la pantalla de entrada.
        if "correo" not in session:
            return redirect(url_for("login"))
        return vista(*args, **kwargs)
    return envoltura


# FUNCIÓN decoradora: además de la sesión, exige que el rol sea "admin".
def solo_admin(vista):
    @wraps(vista)
    def envoltura(*args, **kwargs):
        # CONDICIONAL: primero reviso que haya sesión; si no, va a entrar.
        if "correo" not in session:
            return redirect(url_for("login"))
        # CONDICIONAL: con sesión pero sin rol de administrador, respondo 403 (prohibido).
        if session.get("rol") != "admin":
            abort(403)
        return vista(*args, **kwargs)
    return envoltura


# Manejador del error 403: en vez del error seco, muestro una página amable.
@app.errorhandler(403)
def sin_permiso(error):
    return render_template("sin_permiso.html"), 403


# Manejador del error 413: la petición completa pasa del límite permitido.
@app.errorhandler(413)
def peticion_muy_grande(error):
    # CONDICIONAL: sin sesión, primero mando a entrar.
    if "correo" not in session:
        return redirect(url_for("login"))
    mensaje = ("La carga es demasiado grande. Vercel rechaza peticiones de más de 4.5 MB; "
               "sube archivos más livianos (máximo 4 MB cada uno).")
    return tablero_fabrica(mensaje), 413


# FUNCIÓN: calcula el promedio de una lista de números.
def calcular_promedio(valores):
    # VARIABLE de tipo float: acumulador donde sumo uno por uno.
    total = 0.0
    # BUCLE: recorro cada valor de la lista.
    for v in valores:
        total = total + float(v)
    # CONDICIONAL: si la lista está vacía, devuelvo 0 para no dividir entre cero.
    if len(valores) == 0:
        return 0.0
    return total / len(valores)


# FUNCIÓN: decide el certificado según AGENTS.md (límites INCLUYEN el valor).
def decidir_certificado(promedio, asistencia):
    # CONDICIONAL 1: la asistencia manda; si es menor a 80 no hay certificado.
    if asistencia < 80:
        return "Sin certificado"
    # CONDICIONAL 2: con asistencia suficiente, el promedio decide.
    elif promedio >= 70:
        return "Aprobación"
    # CONDICIONAL 3 (caso restante): buena asistencia pero nota menor a 70.
    else:
        return "Participación"


# FUNCIÓN: lee los Excel y deja todo calculado para las páginas.
# Los parámetros son opcionales: si no me pasan archivos, uso los de la carpeta Insumos/.
def cargar_datos(maestro=None, evaluaciones=None):
    # CONDICIONAL: sin archivo maestro, uso la ruta de fábrica.
    if maestro is None:
        maestro = RUTA_MAESTRO
    # CONDICIONAL: sin archivo de evaluaciones, uso la ruta de fábrica.
    if evaluaciones is None:
        evaluaciones = RUTA_EVALS

    # --- Maestro: lista oficial (VARIABLE de tipo lista de dicts). ---
    # FUNCIÓN: openpyxl acepta tanto una ruta del disco como un archivo en memoria.
    wb_m = openpyxl.load_workbook(maestro, data_only=True)
    ws_m = wb_m["Estudiantes"]
    maestro = {}  # VARIABLE de tipo dict: clave (id, programa) -> datos del estudiante.
    # BUCLE: salto el encabezado (fila 1) y leo cada fila.
    for fila in ws_m.iter_rows(min_row=2, values_only=True):
        sid = str(fila[0])  # VARIABLE de tipo str: el ID siempre como texto.
        maestro[(sid, fila[3])] = {
            "id": sid, "nombre": fila[1], "correo": fila[2],
            "programa": fila[3], "cohorte": str(fila[4]),
        }

    # --- Evaluaciones agrupadas por Identificacion + Programa. ---
    wb_e = openpyxl.load_workbook(evaluaciones, data_only=True)
    ws_e = wb_e["Evaluaciones"]
    grupos = defaultdict(list)  # VARIABLE de tipo dict: clave (id, programa) -> lista de filas.
    # BUCLE: junto cada fila de nota/asistencia en su grupo.
    for fila in ws_e.iter_rows(min_row=2, values_only=True):
        sid = str(fila[0])
        grupos[(sid, fila[1])].append(fila)

    # --- Cruce y cálculo (un grupo = un posible certificado). ---
    resultados = []  # VARIABLE de tipo lista: lo que verá la tabla.
    # BUCLE: proceso cada grupo con las dos FUNCIONES de arriba.
    for (sid, programa), filas in grupos.items():
        # CONDICIONAL: el huérfano (sin maestro) no entra a la tabla, va a inconsistencias.
        if (sid, programa) not in maestro:
            continue
        notas = [float(f[3]) for f in filas]  # VARIABLE de tipo lista de float.
        asists = [float(f[4]) for f in filas]
        promedio = calcular_promedio(notas)  # FUNCIÓN.
        asistencia = calcular_promedio(asists)  # FUNCIÓN reutilizada.
        tipo = decidir_certificado(promedio, asistencia)  # FUNCIÓN con CONDICIONALES.
        resultados.append({
            "id": sid, "nombre": maestro[(sid, programa)]["nombre"],
            "programa": programa, "modulos": len(filas),
            "promedio": round(promedio, 2), "asistencia": round(asistencia, 2),
            "tipo": tipo, "cohorte": maestro[(sid, programa)]["cohorte"],
        })
    # Orden para la tabla: programa y nombre.
    resultados.sort(key=lambda r: (r["programa"], r["nombre"]))

    # --- Contadores de arriba (VARIABLES de tipo int). ---
    contadores = {"Aprobación": 0, "Participación": 0, "Sin certificado": 0}
    # BUCLE con CONDICIONALES: cuento cada tipo.
    for r in resultados:
        if r["tipo"] == "Aprobación":
            contadores["Aprobación"] = contadores["Aprobación"] + 1
        elif r["tipo"] == "Participación":
            contadores["Participación"] = contadores["Participación"] + 1
        else:
            contadores["Sin certificado"] = contadores["Sin certificado"] + 1

    # --- Inconsistencias (las 3 detectadas). ---
    sin_evaluacion = [m for clave, m in maestro.items() if clave not in grupos]
    huerfanos = []  # evaluaciones sin maestro.
    # BUCLE: reviso cada grupo buscando al que no tiene ficha en el maestro.
    for (sid, programa), filas in grupos.items():
        # CONDICIONAL: si no está en el maestro, lo guardo como huérfano.
        if (sid, programa) not in maestro:
            notas_h = [float(f[3]) for f in filas]
            huerfanos.append({"id": sid, "programa": programa,
                              "filas": len(filas),
                              "promedio": round(calcular_promedio(notas_h), 2)})
    con_ceros = []  # caso con notas en 0 para revisión (posible abandono).
    # BUCLE con CONDICIONAL: busco resultados con algún 0 en sus notas.
    for (sid, programa), filas in grupos.items():
        # CONDICIONAL: solo miro grupos válidos (con maestro).
        if (sid, programa) in maestro:
            notas = [float(f[3]) for f in filas]
            tiene_cero = False
            # BUCLE: reviso nota por nota.
            for n in notas:
                # CONDICIONAL: marco si alguna nota es cero.
                if n == 0:
                    tiene_cero = True
            if tiene_cero:
                con_ceros.append({"id": sid,
                                  "nombre": maestro[(sid, programa)]["nombre"],
                                  "programa": programa})
    inconsistencias = {"sin_evaluacion": sin_evaluacion,
                       "huerfanos": huerfanos, "con_ceros": con_ceros}

    programas = sorted(set(m["programa"] for m in maestro.values()))
    return resultados, contadores, inconsistencias, programas


# VARIABLE de tipo tupla: datos listos una sola vez al arrancar.
RESULTADOS, CONTADORES, INCONSISTENCIAS, PROGRAMAS = cargar_datos()


# FUNCIÓN: arma un Excel en memoria con los resultados ya calculados.
def construir_excel(resultados):
    # VARIABLE de tipo Workbook: libro de Excel nuevo, todavía vacío.
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Resultados"
    # VARIABLE de tipo lista: los títulos de las columnas del archivo.
    encabezados = ["Identificacion", "Nombre_Completo", "Programa", "Cohorte",
                   "Modulos", "Promedio", "Asistencia_Pct", "Certificado"]
    hoja.append(encabezados)
    # BUCLE: una fila por cada resultado.
    for r in resultados:
        hoja.append([r["id"], r["nombre"], r["programa"], r["cohorte"],
                     r["modulos"], r["promedio"], r["asistencia"], r["tipo"]])
    # BytesIO es como una caja de cartón dentro de la memoria: guardo el archivo ahí
    # como si fuera una memoria USB, pero sin tocar el disco de la computadora.
    memoria = io.BytesIO()
    libro.save(memoria)  # FUNCIÓN: escribe el Excel dentro de esa caja.
    memoria.seek(0)  # FUNCIÓN: rebobino al inicio para poder entregarlo.
    return memoria


# FUNCIÓN: decide qué inconsistencias puede ver la sesión actual.
def inconsistencias_visibles():
    # CONDICIONAL: solo el administrador ve los avisos del negocio.
    if session.get("rol") == "admin":
        return INCONSISTENCIAS
    return {"sin_evaluacion": [], "huerfanos": [], "con_ceros": []}


# FUNCIÓN: dibuja el tablero con los datos que le pase.
def render_tablero(resultados, contadores, programas, elegido="Todos",
                   origen=ORIGEN_FABRICA, datos_subidos=False,
                   error=None, inconsistencias=None):
    # CONDICIONAL: si no me pasan inconsistencias, uso las que corresponden a la sesión.
    if inconsistencias is None:
        inconsistencias = inconsistencias_visibles()
    return render_template("index.html", tabla=resultados, contadores=contadores,
                           inconsistencias=inconsistencias, programas=programas,
                           elegido=elegido, origen=origen,
                           datos_subidos=datos_subidos, error=error)


# FUNCIÓN: dibuja el tablero con los datos de fábrica (y un error si hace falta).
def tablero_fabrica(error=None):
    return render_tablero(RESULTADOS, CONTADORES, PROGRAMAS, error=error)


# Ruta de entrada: muestra el formulario y valida correo y contraseña.
@app.route("/login", methods=["GET", "POST"])
def login():
    # VARIABLE de tipo str o None: guarda el mensaje rojo cuando las credenciales fallan.
    error = None
    # CONDICIONAL: si ya hay sesión iniciada, no tiene sentido volver a entrar.
    if "correo" in session:
        return redirect(url_for("index"))
    # CONDICIONAL: solo valido cuando el formulario se envía (POST).
    if request.method == "POST":
        correo = request.form.get("correo", "").strip()  # VARIABLE de tipo str.
        clave = request.form.get("clave", "")  # VARIABLE de tipo str.
        usuario = verificar_credenciales(correo, clave)  # FUNCIÓN.
        # CONDICIONAL: si no coincide ningún usuario, muestro el error.
        if usuario is None:
            error = "Correo o contraseña incorrectos."
        else:
            # VARIABLES de sesión: guardo quién entró y su rol para el resto de páginas.
            session["correo"] = correo
            session["rol"] = usuario["rol"]
            return redirect(url_for("index"))
    return render_template("login.html", error=error, modo_demo=MODO_DEMO,
                           demo_admin=(ADMIN_CORREO, ADMIN_CLAVE),
                           demo_usuario=(USUARIO_CORREO, USUARIO_CLAVE))


# Ruta de salida: borra la sesión y devuelve a la pantalla de entrada.
@app.route("/salir")
def salir():
    session.clear()  # FUNCIÓN: elimina todo lo guardado en la sesión.
    return redirect(url_for("login"))


# Ruta del panel: recibe los dos Excel EN MEMORIA y recalcula los resultados.
@app.route("/cargar", methods=["GET", "POST"])
@solo_admin
def cargar():
    # CONDICIONAL: entrar por GET solo lleva al tablero; el panel vive en la página principal.
    if request.method == "GET":
        return redirect(url_for("index"))

    # VARIABLES de tipo FileStorage: los dos archivos que llegaron del formulario.
    archivo_maestro = request.files.get("maestro")
    archivo_evaluaciones = request.files.get("evaluaciones")
    # VARIABLE de tipo str: qué botón se pulsó ("ver" o "excel").
    accion = request.form.get("accion", "ver")

    # CONDICIONAL 1: falta alguno de los dos archivos.
    if (archivo_maestro is None or archivo_evaluaciones is None or
            archivo_maestro.filename == "" or archivo_evaluaciones.filename == ""):
        return tablero_fabrica(
            "Debes elegir los dos archivos: maestro de estudiantes y registro de evaluaciones.")

    # BUCLE con CONDICIONAL 2: los dos deben ser Excel modernos (.xlsx).
    for archivo in (archivo_maestro, archivo_evaluaciones):
        if not archivo.filename.lower().endswith(".xlsx"):
            return tablero_fabrica("Los archivos deben ser Excel con extensión .xlsx.")

    # FUNCIÓN: leo cada archivo a la memoria, sin escribirlo en el disco.
    datos_maestro = archivo_maestro.read()
    datos_evaluaciones = archivo_evaluaciones.read()

    # CONDICIONAL 3: algún archivo pasa de 4 MB.
    if len(datos_maestro) > LIMITE_ARCHIVO or len(datos_evaluaciones) > LIMITE_ARCHIVO:
        return tablero_fabrica(
            "Cada archivo debe pesar menos de 4 MB. Reduce el Excel y vuelve a intentarlo.")

    # CONDICIONAL 4: alguno llegó vacío.
    if len(datos_maestro) == 0 or len(datos_evaluaciones) == 0:
        return tablero_fabrica("Alguno de los archivos está vacío. Revisa que tenga datos.")

    try:
        # FUNCIÓN: convierto los bytes en una "caja" legible y reutilizo cargar_datos.
        # (No duplico la lógica: cargar_datos ya sabe promediar y clasificar.)
        resultados, contadores, inconsistencias, programas = cargar_datos(
            io.BytesIO(datos_maestro), io.BytesIO(datos_evaluaciones))
    except KeyError:
        # CONDICIONAL: al Excel le falta la hoja "Estudiantes" o "Evaluaciones".
        return tablero_fabrica(
            "El Excel no tiene las hojas esperadas: 'Estudiantes' y 'Evaluaciones'.")
    except Exception:
        # CONDICIONAL: cualquier otro problema al leer (archivo dañado, columnas raras...).
        return tablero_fabrica(
            "No pudimos leer los archivos. Revisa que sean Excel válidos y que no estén dañados.")

    # VARIABLE de tipo str: nombre de los archivos, para mostrar el origen en pantalla.
    origen = f"Archivos subidos: {archivo_maestro.filename} y {archivo_evaluaciones.filename}"

    # CONDICIONAL: si pidió Excel, armo el archivo en memoria y lo devuelvo.
    if accion == "excel":
        memoria = construir_excel(resultados)
        return send_file(
            memoria,
            as_attachment=True,
            download_name="resultados_certificados.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    # Si no, muestro la tabla con los datos recién calculados.
    return render_tablero(resultados, contadores, programas, origen=origen,
                          datos_subidos=True, inconsistencias=inconsistencias)


# Ruta de descarga: entrega un Excel con los resultados de fábrica.
@app.route("/descargar")
@solo_admin
def descargar():
    memoria = construir_excel(RESULTADOS)
    return send_file(
        memoria,
        as_attachment=True,
        download_name="resultados_fabrica.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# Ruta principal: contadores + filtro + tabla + inconsistencias.
@app.route("/")
@requiere_login
def index():
    # VARIABLE de tipo str: lo que eligió el filtro (o "Todos").
    elegido = request.args.get("programa", "Todos")
    # CONDICIONAL: si eligió un programa, filtro la tabla; si no, muestro todo.
    if elegido == "Todos":
        tabla = RESULTADOS
    else:
        tabla = [r for r in RESULTADOS if r["programa"] == elegido]
    # FUNCIÓN: dibujo el tablero con los datos de fábrica.
    return render_tablero(tabla, CONTADORES, PROGRAMAS, elegido=elegido)


# FUNCIÓN: busca un resultado por identificación y programa.
def buscar_resultado(sid, programa):
    # VARIABLE de tipo dict o None: aquí guardo lo encontrado.
    encontrado = None
    # BUCLE: recorro los resultados ya calculados.
    for r in RESULTADOS:
        # CONDICIONAL: comparo ID (texto) y programa.
        if r["id"] == str(sid) and r["programa"] == programa:
            encontrado = r
    return encontrado


# Ruta detalle: un certificado por estudiante (solo muestra, no calcula).
@app.route("/certificado/<sid>")
@requiere_login
def certificado(sid):
    # VARIABLE de tipo lista: el certificado se dibuja igual para uno o varios.
    elegidos = [r for r in RESULTADOS if r["id"] == str(sid)]
    # CONDICIONAL: si no existe, muestro la página de aviso.
    if len(elegidos) == 0:
        return render_template("sin_certificado.html", sid=sid)
    return render_template("certificado.html", certificados=elegidos,
                           hoy=date.today().strftime("%d/%m/%Y"))


# Ruta de emisión: recibe las casillas marcadas y genera sus certificados.
@app.route("/generar", methods=["POST"])
@solo_admin
def generar():
    # VARIABLE de tipo lista de str: cada casilla llega como "id|programa".
    seleccion = request.form.getlist("seleccion")
    elegidos = []  # VARIABLE de tipo lista: certificados que sí se pueden emitir.
    # BUCLE: reviso cada casilla marcada.
    for clave in seleccion:
        # CONDICIONAL: protejo el formato por si llega algo raro.
        if "|" not in clave:
            continue
        sid, programa = clave.split("|", 1)  # VARIABLES de tipo str.
        r = buscar_resultado(sid, programa)  # FUNCIÓN.
        # CONDICIONAL: solo emito si existe y NO es "Sin certificado".
        if r is not None and r["tipo"] != "Sin certificado":
            elegidos.append(r)
    # CONDICIONAL: si no marcó nada válido, aviso en pantalla.
    if len(elegidos) == 0:
        return render_template("sin_certificado.html", sid="ninguna selección")
    return render_template("certificado.html", certificados=elegidos,
                           hoy=date.today().strftime("%d/%m/%Y"))


if __name__ == "__main__":
    # Arranco el salón en http://localhost:5000
    app.run(debug=True)
